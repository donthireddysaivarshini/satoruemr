"""validate_exports.py - offline validation of the new Excel export scopes.

Runs build_workbook('camp-workbook=...' / 'participant-report=...') against
synthetic camps/participants/reports (no network) and asserts the required
structure and human-readable content. Also runs a network-free sanity pass of
scale_export.row alignment for all four scales.
"""

import io
import json
import os
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

import openpyxl

import export_emr
import scale_export

FAILURES = []
CHECKS = [0]


def check(name, cond, extra=''):
    CHECKS[0] += 1
    if not cond:
        FAILURES.append(f'{name}{": " + str(extra) if extra else ""}')
        print(f'  FAIL {name} {extra}')
    else:
        print(f'  ok   {name}')


# ---------------------------------------------------------------------------
# synthetic data
# ---------------------------------------------------------------------------

PART_A = {
    '_id': 'participant-a', 'type': 'contact', 'contact_type': 'participant',
    'name': 'Kamala', 'cr_no': 'CR001', 'age': 1, 'gender': 'female',
    'phone': '+919000000001', 'hospital_name': 'NIMS Hospital',
    'ward_dept_no': 'W-10', 'address': 'Hyderabad', 'education_level': None,
    'occupation': None, 'referred_by': 'Dr Rao',
    'referral_department': 'Pediatrics', 'referral_complaint': 'Delay',
    'notes': 'repeat visitor', 'registration_date': 1756339200000,
    'parent': {'_id': 'camp-1'},
}
PART_B = {
    '_id': 'participant-b', 'type': 'contact', 'contact_type': 'participant',
    'name': 'Ravi', 'cr_no': 'CR002', 'age': 68, 'gender': 'male',
    'phone': '+919000000002', 'hospital_name': 'NIMS Hospital',
    'parent': {'_id': 'camp-1'},
}
PART_C = {  # other camp: must not leak into camp-1 exports
    '_id': 'participant-c', 'type': 'contact', 'contact_type': 'participant',
    'name': 'Other', 'cr_no': 'CR003', 'parent': {'_id': 'camp-2'},
}
CAMP_1 = {'_id': 'camp-1', 'type': 'contact', 'contact_type': 'screening_camp',
          'name': 'NIMS', 'notes': 'main camp'}
CAMP_2 = {'_id': 'camp-2', 'type': 'contact', 'contact_type': 'screening_camp',
          'name': 'OtherCamp'}
SETTINGS = {'_id': 'settings', 'settings': {'contact_types': [
    {'id': 'screening_camp', 'person': False},
    {'id': 'participant', 'person': True},
]}}
SUBMITTER = {'_id': 'user-contact-1', 'type': 'contact',
             'contact_type': 'person', 'name': 'Varshini',
             'phone': '+919000000009'}


def report(form, sid, fields, reported_date):
    return {
        '_id': f'report-{form}-{sid}', 'type': 'data_record', 'form': form,
        'content_type': 'xml', 'reported_date': reported_date,
        'contact': SUBMITTER,  # submitter; subject is fields.patient_uuid
        'fields': dict(fields, patient_uuid=sid),
        'hidden_fields': [],
    }


# Vineland: fields taken from the real form structure
VINELAND_FIELDS = {
    'patient_id': 'CR001', 'patient_name': 'Kamala',
    'g_child': {'assessment_date': '2026-08-29', 'child_dob': '2025-07-17',
                'ca_days': 408, 'ca_months_dec': 13.4,
                'ca_months_floor': 13, 'ca_years_display': 1,
                'ca_rem_months': 1, 'ca_band': '0-1', 'age_valid': 'yes',
                'show_all': 'yes', 'ca_display': '1 year 1 month'},
    'g_y0_1': {'item_1': 'yes', 'item_2': 'yes', 'item_3': 'no'},
    'item_1_score': 1, 'item_2_score': 1, 'item_3_score': 0,
    'raw_score': 22, 'total_yes_items': 22, 'total_no_items': 3,
    'total_not_assessed_items': 1, 'social_age_months': 22,
    'social_age_years': 1, 'social_age_remaining_months': 10,
    'social_quotient_raw': 164.12, 'social_quotient_rounded': 164,
    'sq_interpretation': 'Average', 'na_reason': '', 'clinical_notes': 'Good eye contact.',
    'scoring_notes': 'Scored per 0-1 band.',
    'g_supporting_docs': {'support_doc_1': 'assessment.pdf'},
}

# DDST: uses g_meta
DDST_FIELDS = {
    'patient_id': 'CR001', 'patient_name': 'Kamala',
    'g_meta': {'test_date': '2026-08-29', 'child_dob': '2025-07-17',
               'ca_days': 408, 'ca_months_dec': 13.4, 'ca_months_floor': 13,
               'ca_band': '0-1', 'show_all': 'yes', 'ca_note': 'x',
               'dob_warning': ''},
    'g_ps': {'ps_01': 'P', 'ps_02': 'F', 'ps_03': 'R'},
    'g_fm': {'fm_01': 'P', 'fm_02': 'R'},
    'd_ps_status': 'Normal', 'd_fm_status': 'Delayed',
    'd_lang_status': 'Normal', 'd_gm_status': 'Normal',
    'caution_domains': '', 'any_suspect': 'no', 'overall': 'Normal',
    'result_note': 'composed text duplicated by results rows',
}


def read_only_ws(wb, title):
    ws = wb[title]
    rows = list(ws.iter_rows(values_only=True))
    return rows[0], rows[1:]


# ---------------------------------------------------------------------------
# 1. camp-workbook scope
# ---------------------------------------------------------------------------

def test_camp_workbook():
    print('== camp-workbook scope ==')
    docs = [SETTINGS, CAMP_1, CAMP_2, PART_A, PART_B, PART_C, SUBMITTER,
            report('vineland', 'participant-a', VINELAND_FIELDS, 1756425600000),
            report('ddst', 'participant-a', DDST_FIELDS, 1756425700000),
            report('ddst', 'participant-b', dict(DDST_FIELDS, patient_name='Ravi'), 1756425800000)]
    wb = export_emr.build_workbook('camp-workbook=camp-1',
                                   FakeDb(docs))

    titles = [ws.title for ws in wb.worksheets]
    check('sheets: Participants + only completed scales',
          titles == ['Participants', 'DDST-II', 'VSMS'], titles)

    # --- Participants sheet: one row per camp member ---
    hdr, rows = read_only_ws(wb, 'Participants')
    check('Participants: 2 rows (only camp-1 members)', len(rows) == 2, len(rows))
    idx = {h: i for i, h in enumerate(hdr)}
    row_a = next(r for r in rows if r[idx['Name']] == 'Kamala')
    row_b = next(r for r in rows if r[idx['Name']] == 'Ravi')
    check('Participants: scales completed listed',
          row_a[idx['Scales completed']] == 'DDST-II, VSMS'
          and row_b[idx['Scales completed']] == 'DDST-II',
          (row_a[idx['Scales completed']], row_b[idx['Scales completed']]))
    check('Participants: gender human-readable',
          row_a[idx['Gender']] == 'Female', row_a[idx['Gender']])
    check('Participants: no internal _id column',
          '_id' not in idx and 'Patient ID' not in hdr)

    # --- DDST-II sheet: one row per completed assessment ---
    hdr, rows = read_only_ws(wb, 'DDST-II')
    check('DDST-II: header + 2 rows', len(rows) == 2, len(rows))
    idx = {h: i for i, h in enumerate(hdr)}
    check('DDST-II: no technical keys in header',
          not any(h and ('report.' in h or 'ca_days' in h or 'patient_uuid' in h)
                  for h in hdr))
    row_a = next(r for r in rows if r[idx['Participant']] == 'Kamala')
    row_b = next(r for r in rows if r[idx['Participant']] == 'Ravi')
    # P -> Pass resolved via the form's own choice list
    check('DDST-II: choice P resolved to label',
          any(isinstance(v, str) and v == 'Pass' for v in row_a), row_a)
    check('DDST-II: no raw P/F values remain',
          not any(v in ('P', 'F', 'NO', 'R') for v in row_a))
    check('DDST-II: submitted by from contact name',
          'Varshini' in (row_a[idx['Submitted by']] or ''), row_a[idx['Submitted by']])
    check('DDST-II: DOB filled from form', row_a[idx['DOB']] == '2025-07-17',
          row_a[idx['DOB']])
    check('DDST-II: chronological age joined',
          any(isinstance(v, str) and 'year' in v for v in row_a))
    check('DDST-II: no raw calculated fields exposed',
          not any(v == 'x' for v in row_a))  # g_meta.ca_note value 'x' hidden

    # --- VSMS sheet: one row (only Kamala completed it) ---
    hdr, rows = read_only_ws(wb, 'VSMS')
    check('VSMS: header + 1 row (Ravi absent)', len(rows) == 1, len(rows))
    idx = {h: i for i, h in enumerate(hdr)}
    row = rows[0]
    check('VSMS: question label from form',
          '1. Cries, laughs' in hdr and '1. Cries, laughs - score' in hdr,
          [h for h in hdr if 'Cries' in str(h)])
    qcol = hdr.index('1. Cries, laughs')
    check('VSMS: answer resolved yes -> Yes', row[qcol] == 'Yes', row[qcol])
    score_col = hdr.index('1. Cries, laughs - score')
    check('VSMS: item score column joined', row[score_col] == 1, row[score_col])
    res_idx = {h: i for i, h in enumerate(hdr)}
    check('VSMS: SQ interpretation present',
          res_idx.get('SQ Interpretation') is not None
          and row[res_idx['SQ Interpretation']] == 'Average',
          row[res_idx['SQ Interpretation']] if 'SQ Interpretation' in res_idx else 'missing')
    check('VSMS: raw calc fields absent from header',
          not any('ca_months_floor' in str(h) for h in hdr))


# ---------------------------------------------------------------------------
# 2. participant-report scope
# ---------------------------------------------------------------------------

def test_participant_report():
    print('== participant-report scope ==')
    docs = [SETTINGS, CAMP_1, PART_A, PART_B, SUBMITTER,
            report('vineland', 'participant-a', VINELAND_FIELDS, 1756425600000),
            report('ddst', 'participant-a', DDST_FIELDS, 1756425700000)]
    wb = export_emr.build_workbook('participant-report=participant-a',
                                   FakeDb(docs))

    titles = [ws.title for ws in wb.worksheets]
    check('sheets: Participant + completed scales only',
          titles == ['Participant', 'DDST-II', 'VSMS'], titles)

    hdr, rows = read_only_ws(wb, 'Participant')
    idx = {h: i for i, h in enumerate(hdr)}
    kv = {r[0]: r[1] for r in rows}
    check('Participant: key/value layout with Field header',
          hdr[0] == 'Field' and hdr[1] == 'Value')
    check('Participant: human fields present',
          kv.get('Participant ID') == 'CR001' and kv.get('Name') == 'Kamala'
          and kv.get('Camp') == 'NIMS', kv.get('Participant ID'))
    check('Participant: scales completed listed',
          kv.get('Scales completed') == 'DDST-II, VSMS', kv.get('Scales completed'))
    check('Participant: empty fields omitted',
          all(v not in (None, '') for v in kv.values()))

    # VSMS sectioned sheet
    ws = wb['VSMS']
    rows = list(ws.iter_rows(values_only=True))
    text = [tuple(r) for r in rows]
    flat_cells = [c for r in rows for c in r if c not in (None, '')]
    check('VSMS: scale title section',
          any(isinstance(c, str) and 'VSMS' in c or 'Vineland' in c for c in flat_cells))
    check('VSMS: question row with answer + score',
          any(isinstance(r, tuple) and r[0] == '1. Cries, laughs'
              and r[1] == 'Yes' and r[2] == 1 for r in text),
          [r for r in text if r and r[0] and 'Cries' in str(r[0])])
    check('VSMS: scoring results section present',
          any(r and r[0] == 'Scoring results' for r in text))
    check('VSMS: interpretation row',
          any(r and r[0] == 'SQ Interpretation' and r[1] == 'Average'
              for r in text))
    check('VSMS: chronological age joined row',
          any(r and r[0] == 'Chronological age' for r in text))
    check('VSMS: no internal fields in output',
          not any(isinstance(c, str) and ('ca_months_floor' in c or 'patient_uuid' in c)
                  for c in flat_cells))
    check('VSMS: notes section with clinical notes',
          any(r and r[0] == 'Clinical notes' for r in text))
    check('VSMS: supporting documents listed',
          any(isinstance(c, str) and 'assessment.pdf' in c for c in flat_cells))

    # DDST sectioned sheet
    rows = list(wb['DDST-II'].iter_rows(values_only=True))
    flat_cells = [c for r in rows for c in r if c not in (None, '')]
    check('DDST-II: domain sections grouped (Personal-Social band label)',
          any(isinstance(c, str) and 'Personal' in c for c in flat_cells),
          [c for c in flat_cells if isinstance(c, str) and 'Personal' in c][:2])
    check('DDST-II: pass label resolved in sections',
          any(c == 'Pass' for c in flat_cells))
    check('DDST-II: overall result row',
          any(r and r[0] == 'Overall' for r in rows),
          [r[0] for r in rows if r and r[0]][:8])

    # scope isolation: this docs fixture has no reports for Ravi at all,
    # so his workbook legitimately contains only the Participant sheet
    wb2 = export_emr.build_workbook('participant-report=participant-b', FakeDb(docs))
    titles2 = [ws.title for ws in wb2.worksheets]
    check('participant-report: Ravi (no reports) gets Participant sheet only',
          titles2 == ['Participant'], titles2)
    docs_ravi = docs + [report('ddst', 'participant-b',
                               dict(DDST_FIELDS, patient_name='Ravi'), 1756425800000)]
    wb3 = export_emr.build_workbook('participant-report=participant-b', FakeDb(docs_ravi))
    titles3 = [ws.title for ws in wb3.worksheets]
    check('participant-report: Ravi with DDST gets exactly Participant + DDST-II',
          titles3 == ['Participant', 'DDST-II'], titles3)


# ---------------------------------------------------------------------------
# 3. scale_export header/row alignment for all scales
# ---------------------------------------------------------------------------

def test_scale_row_alignment():
    print('== scale_row / scale_headers alignment ==')
    samples = {
        'vineland': (VINELAND_FIELDS, PART_A, 'Varshini', ['assessment.pdf']),
        'ddst': (DDST_FIELDS, PART_A, 'Varshini', []),
    }
    samples['dst'] = ({
        'patient_id': 'CR001', 'patient_name': 'Kamala',
        'g_child': {'assessment_date': '2026-08-29', 'child_dob': '2025-07-17',
                    'ca_days': 408, 'ca_months_dec': 13.4, 'ca_band': '0-1',
                    'band_1_passed': 5, 'band_1_failed': 1,
                    'band_1_not_assessed': 0, 'band_1_earned': 5},
        'g_band_1': {'item_1': 'pass', 'item_2': 'fail', 'item_3': 'not_assessed'},
        'total_passed_items': 12, 'developmental_age_months': 15,
        'developmental_quotient_rounded': 100,
    }, PART_A, 'Varshini', [])
    samples['moca_assessment'] = ({
        'g_admin': {'assessment_datetime': '2026-08-29T10:30:00.000Z'},
        'g_visuospatial': {'trail_making': 'yes'},
        'moca_visuospatial_total': 1, 'moca_raw_total': 28,
        'moca_education_adj': 0, 'moca_total_score': 28,
        'moca_interpretation': 'Normal',
    }, PART_B, 'Varshini', [])

    for form, (fields, part, submitter, atts) in samples.items():
        headers = scale_export.scale_headers(form)
        row = scale_export.scale_row(form, fields, part, submitter, atts)
        check(f'{form}: headers/row same length',
              len(headers) == len(row), (len(headers), len(row)))
        # sections alignment
        sections = scale_export.scale_sections(form, fields, part, submitter, atts)
        check(f'{form}: sections non-empty', len(sections) >= 3, len(sections))
        # no technical keys as labels
        bad = [s for s, rows_ in sections for r in rows_
               if isinstance(r[0], str) and (r[0].startswith('report.')
                                            or 'patient_uuid' in r[0]
                                            or 'ca_months_floor' in r[0])]
        check(f'{form}: no technical labels in sections', not bad, bad[:3])


class FakeDb:
    """Minimal stand-in for EmrDb (fetch_all_docs only)."""

    def __init__(self, docs):
        self._docs = docs

    def fetch_all_docs(self):
        return self._docs

    def fetch_doc(self, doc_id):
        for d in self._docs:
            if d.get('_id') == doc_id:
                return d
        return None


def main():
    test_camp_workbook()
    test_participant_report()
    test_scale_row_alignment()
    print(f'\n{CHECKS[0]} checks, {len(FAILURES)} failures')
    if FAILURES:
        print('FAILURES:')
        for f in FAILURES:
            print(f'  - {f}')
        sys.exit(1)
    print('ALL EXPORT VALIDATION PASSED')


if __name__ == '__main__':
    main()
