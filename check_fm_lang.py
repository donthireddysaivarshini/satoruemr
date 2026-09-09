import openpyxl
import re

wb = openpyxl.load_workbook(r'C:\Projects\EMRCHT\satoru-config\forms\app\ddst.xlsx')
ws = wb['survey']

# Check FM items (rows 57-78)
print('=== Fine Motor Items relevant expressions ===')
for row_idx in range(57, 80):
    relevant = ws.cell(row=row_idx, column=7).value
    name = ws.cell(row=row_idx, column=3).value
    if name or relevant:
        rel_str = str(relevant) if relevant else ''
        for ch in ['\u25cb', '\u25a1']:
            rel_str = rel_str.replace(ch, '<square>')
        rel_str = rel_str[:50]
        age_match = re.search(r'>= ([\d.]+)', rel_str)
        age_threshold = float(age_match.group(1)) if age_match else None
        print(f'Row {row_idx}: {rel_str}... -> age={age_threshold}')

print()

# Check Lang items (rows 89-127)
print('=== Language Items relevant expressions ===')
for row_idx in range(89, 130):
    relevant = ws.cell(row=row_idx, column=7).value
    name = ws.cell(row=row_idx, column=3).value
    if name or relevant:
        rel_str = str(relevant) if relevant else ''
        for ch in ['\u25cb', '\u25a1']:
            rel_str = rel_str.replace(ch, '<square>')
        rel_str = rel_str[:50]
        age_match = re.search(r'>= ([\d.]+)', rel_str)
        age_threshold = float(age_match.group(1)) if age_match else None
        print(f'Row {row_idx}: {rel_str}... -> age={age_threshold}')