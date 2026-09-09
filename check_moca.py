import openpyxl
wb = openpyxl.load_workbook(r'C:\Projects\EMRCHT\satoru-config\forms\app\moca_assessment.xlsx')
ws = wb['survey']

# Find the calculation rows
print('=== MoCA Calculation Rows ===')
for row_idx in range(1, ws.max_row + 1):
    cell = ws.cell(row=row_idx, column=12)  # calculation column
    if cell.value and 'moca' in str(cell.value).lower():
        rel = ws.cell(row=row_idx, column=7).value  # relevant
        name = ws.cell(row=row_idx, column=2).value  # name
        print(f'Row {row_idx}: name={name}, calc={cell.value}, rel={rel}')

print()

# Also check education adjustment row
print('=== Education Adjustment Row ===')
for row_idx in range(1, ws.max_row + 1):
    cell = ws.cell(row=row_idx, column=12)  # calculation column
    if cell.value and 'moca_education_adj' in str(cell.value):
        rel = ws.cell(row=row_idx, column=7).value  # relevant
        name = ws.cell(row=row_idx, column=2).value  # name
        print(f'Row {row_idx}: name={name}, calc={cell.value}, rel={rel}')

print()

# Check final score row
print('=== Final Score Row ===')
for row_idx in range(1, ws.max_row + 1):
    cell = ws.cell(row=row_idx, column=12)  # calculation column
    if cell.value and 'moca_total_score' in str(cell.value):
        rel = ws.cell(row=row_idx, column=7).value  # relevant
        name = ws.cell(row=row_idx, column=2).value  # name
        print(f'Row {row_idx}: name={name}, calc={cell.value}, rel={rel}')

print()

# Check interpretation row
print('=== Interpretation Row ===')
for row_idx in range(1, ws.max_row + 1):
    cell = ws.cell(row=row_idx, column=12)  # calculation column
    if cell.value and 'moca_interpretation' in str(cell.value):
        rel = ws.cell(row=row_idx, column=7).value  # relevant
        name = ws.cell(row=row_idx, column=2).value  # name
        print(f'Row {row_idx}: name={name}, calc={cell.value}, rel={rel}')