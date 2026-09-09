import openpyxl
wb = openpyxl.load_workbook(r'C:\Projects\EMRCHT\satoru-config\forms\app\moca_assessment.xlsx')
ws = wb['survey']

# Print all rows with calculation
print('=== All rows with calculation ===')
for row_idx in range(1, ws.max_row + 1):
    calc = ws.cell(row=row_idx, column=12).value
    name = ws.cell(row=row_idx, column=2).value
    if calc:
        print(f'Row {row_idx}: name={name}, calc={calc}')