import openpyxl
wb = openpyxl.load_workbook(r'C:\Projects\EMRCHT\satoru-config\forms\app\moca_assessment.xlsx')
ws = wb['survey']

# Print rows 120-131 with column 11 (calculation is column K = 11)
print('=== Rows 120-131, column 11 (calculation) ===')
for row_idx in range(120, 132):
    calc = ws.cell(row=row_idx, column=11).value
    name = ws.cell(row=row_idx, column=2).value
    print(f'Row {row_idx}: name={name}, calc={calc}')

print()

# Also check column 12 and 13
print('=== Rows 120-131, column 12 ===')
for row_idx in range(120, 132):
    calc = ws.cell(row=row_idx, column=12).value
    name = ws.cell(row=row_idx, column=2).value
    print(f'Row {row_idx}: name={name}, calc={calc}')

print()

print('=== Rows 120-131, column 13 ===')
for row_idx in range(120, 132):
    calc = ws.cell(row=row_idx, column=13).value
    name = ws.cell(row=row_idx, column=2).value
    print(f'Row {row_idx}: name={name}, calc={calc}')