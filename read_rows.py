import openpyxl
wb = openpyxl.load_workbook(r'C:\Projects\EMRCHT\satoru-config\forms\app\ddst.xlsx')
ws = wb['survey']

# Print rows 79-130 to find GM items
print('=== Rows 79-130 ===')
for row_idx in range(79, 131):
    relevant = ws.cell(row=row_idx, column=7).value
    name = ws.cell(row=row_idx, column=3).value
    rtype = ws.cell(row=row_idx, column=1).value
    if name or relevant:
        rel_str = str(relevant) if relevant else ''
        name_str = name if name else ''
        # Replace troublesome unicode chars
        for ch in ['\u25cb', '\u25a1']:
            rel_str = rel_str.replace(ch, '<square>')
            name_str = name_str.replace(ch, '<square>')
        rel_str = rel_str[:30]
        name_str = name_str[:20]
        print(f'Row {row_idx}: type={rtype}, relevant={rel_str}..., label={name_str}')