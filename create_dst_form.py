"""Generate forms/app/dst.xlsx - Developmental Screening Test (DST).

Design notes:
- DOB is asked in-form (participant contacts have no dob field) - documented limitation.
- Band visibility: birth through CA band inclusive; future bands hidden; "show all" is trained-user override.
- Scoring at ROOT level guarded by visibility so hidden future bands contribute 0, not failed.
- Band 1 weight uses exact 3 div 7 so all-pass = exactly 180 months.
"""
import openpyxl

wb = openpyxl.Workbook()
ws = wb.active
ws.title = 'survey'
HEADERS = ['type', 'name', 'label', 'hint', 'required', 'appearance', 'relevant',
           'constraint', 'constraint_message', 'calculation', 'default', 'instance::type']
ws.append(HEADERS)
E = ''

def row(**kw):
    # Handle instance::type which is not a valid python kwarg name
    instance_type = kw.pop('instance_type', E)
    ws.append([kw.get(h, E) for h in HEADERS[:-1]] + [instance_type])

# ---------------------------------------------------------------- inputs
row(type='begin group', name='inputs', label='Participant', appearance='field-list',
    relevant="./source = 'user'")
row(type='hidden', name='source')
row(type='hidden', name='source_id')
row(type='begin group', name='user', label='User', appearance='field-list')
row(type='hidden', name='contact_id')
row(type='hidden', name='name')
row(type='end group')
row(type='begin group', name='contact', label='Select child / beneficiary',
    appearance='field-list')
row(type='db:person', name='_id', label='Select participant',
    appearance='select-contact type-participant')
row(type='hidden', name='patient_id')
row(type='hidden', name='name')
row(type='end group')
row(type='end group')

# patient calcs
for n, c in [('patient_uuid', '../inputs/contact/_id'),
             ('patient_id', '../inputs/contact/patient_id'),
             ('patient_name', '../inputs/contact/name'),
             ('screened_by', '../inputs/user/contact_id'),
             ('screened_by_name', '../inputs/user/name')]:
    row(type='calculate', name=n, calculation=c)

# ---------------------------------------------------------------- intro
row(type='begin group', name='g_intro', label='Developmental Screening Test (DST)')
row(type='note', name='dst_disclaimer',
    label="DST is a developmental screening instrument and does not provide a diagnosis. "
          "Results must be interpreted by a qualified clinician. This assessment does not "
          "diagnose developmental, intellectual, neurological or psychiatric conditions.")
row(type='select_one yesno', name='acknowledged', required='yes',
    label='I confirm that this assessment is being completed by or reviewed with a trained/authorized health worker.',
    constraint=". = 'yes'", constraint_message='You must confirm before continuing.')
row(type='end group')

# ---------------------------------------------------------------- child details
row(type='begin group', name='g_child', label='Child details & assessment date')
row(type='date', name='assessment_date', label='Assessment date', required='yes',
    default='today()')
row(type='date', name='child_dob', label='Child date of birth', required='yes',
    hint='Contacts do not store date of birth, so it is captured here.',
    constraint='. <= today() and today() - . <= 5479',
    constraint_message='Date of birth must be today or earlier, and the child must be 15 years or younger.')
row(type='calculate', name='ca_days', calculation='int(${assessment_date} - ${child_dob})')
row(type='calculate', name='ca_months_dec', calculation='${ca_days} div 30.4375')
row(type='calculate', name='ca_months_floor', calculation='floor(${ca_months_dec})')
row(type='calculate', name='age_valid', calculation="if(${ca_days} >= 0 and ${ca_months_dec} <= 180 and ${acknowledged} = 'yes', 'yes', 'no')")
row(type='note', name='dob_warning',
    label="**WARNING:** Child date of birth is missing or invalid. Please enter DOB before scoring. Do not guess age.",
    relevant="${child_dob} = '' or ${age_valid} != 'yes'")
row(type='note', name='ca_display',
    label="Chronological age: **${ca_years_display}y ${ca_rem_months}m** (${ca_days} days).",
    relevant="${age_valid} = 'yes'")
row(type='calculate', name='ca_years_display', calculation='floor(${ca_days} div 365.25)')
row(type='calculate', name='ca_rem_months', calculation='floor((${ca_days} mod 365) div 30.4375)')
row(type='select_one yesno', name='show_all', label='Show all DST items? (trained assessors)',
    default='no', hint='No = show birth through child age band; future bands hidden. Yes = show all bands.')
row(type='end group')

# ---------------------------------------------------------------- bands
BANDS = [
    ('birth_to_3_months',  'Birth to 3 months',   3,   '3 div 7'),
    ('m3_to_6',            '3 to 6 months',       6,   '0.5'),
    ('m6_to_9',            '6 to 9 months',       9,   '0.75'),
    ('m9_to_12',           '9 to 12 months',      12,  '0.75'),
    ('m12_to_18',          '12 to 18 months',     18,  '1.5'),
    ('m18_to_24',          '18 to 24 months',     24,  '1.5'),
    ('m24_to_36',          '2 to 3 years',        36,  '2'),
    ('m36_to_48',          '3 to 4 years',        48,  '2.4'),
    ('m48_to_60',          '4 to 5 years',        60,  '2'),
    ('m60_to_72',          '5 to 6 years',        72,  '2.4'),
    ('m72_to_84',          '6 to 7 years',        84,  '2.4'),
    ('m84_to_96',          '7 to 8 years',        96,  '3'),
    ('m96_to_108',         '8 to 9 years',        108, '2'),
    ('m108_to_120',        '9 to 10 years',       120, '2.4'),
    ('m120_to_132',        '10 to 11 years',      132, '3'),
    ('m132_to_144',        '11 to 12 years',      144, '4'),
    ('m144_to_156',        '12 to 13 years',      156, '2.4'),
    ('m156_to_180',        '13 to 15 years',      181, '4.8'),
]
ITEMS = {
    1: [(n, t, t.endswith('*')) for n, t in [
        (1, 'birth cry present*'), (2, 'equal bilateral movements'), (3, 'responds to bell'),
        (4, 'vocalises sounds*'), (5, 'smiles spontaneously'), (6, 'eyes follow moving objects'),
        (7, 'head steady')]],
}
RAW = [
    "8. reaches for objects|9. laughs aloud*|10. recognizes mother|11. vocalizes for pleasure/babbles*|12. carries objects to mouth|13. rolls over",
    "14. imitates speech sounds*|15. sits by self|16. thumb finger grasp|17. shows curiosity",
    "18. says 3 word, dada, mama etc*|19. stands alone well|20. follows simple instructions*|21. cooperates for dressing",
    "22. many intelligible words*|23. walks, runs well|24. indicates wants|25. scribbles spontaneously",
    "26. says sentences of 2/3 words*|27. points out objects in pictures|28. shows body parts|29. participates in play",
    "30. copies O|31. relates experiences*|32. knows names, uses of common objects|33. begins to ask why?*|34. takes food by self|35. toilet control present",
    "36. buttons up|37. comprehends hunger, cold*|38. plays cooperatively with children|39. repeats 3 digits*|40. tells stories*",
    "41. defines words*|42. makes simple drawings|43. dresses with no supervision|44. describes actions in pictures*|45. gives sensible answers to questions*|46. goes about neighborhood",
    "47. can name primary color|48. plays games governed by rules|49. writes simple words*|50. gains admission to school|51. enjoys constructive play",
    "52. adapts to home, school|53. tells differences of objects|54. spells, reads, writes simple words*|55. enjoys group play|56. knows comparative value of coins",
    "57. combs hair by self|58. makes small purchases|59. competition in school/play|60. tells time",
    "61. tells day, month, year*|62. reads on own initiative*|63. recognizes property rights|64. favourite of fairy tales*|65. muscle coordination games (marbles)|66. bathes self unaided",
    "67. cooperates keenly with companions|68. has various hobbies, collections|69. goes about town freely|70. sex difference in play becomes marked|71. can stay away from home",
    "72. writes occasional short letters|73. comprehends social situations|74. physical feats liked|75. able to discuss problems*",
    "76. enjoys books, newspapers, magazines|77. more independent in spending|78. capable of self-criticism",
    "79. shows foresight, planning, judgment|80. learns from experiences|81. plays difficult games|82. interested in dressing up|83. understands abstract ideas (justice)",
    "84. makes sensible plans for future (job)|85. follows current events*|86. buys own clothing|87. systematises own work|88. purchases for others",
]

def parse_item(txt):
    num, _, rest = txt.partition('. ')
    speech = rest.endswith('*')
    return int(num), rest.rstrip('*').strip(), speech

ALL_BAND_ITEMS = {1: ITEMS[1]}
for i, raw in enumerate(RAW, start=2):
    ALL_BAND_ITEMS[i] = [parse_item(t) for t in raw.split('|')]

def vb_expr(k):
    # Show birth through CA band; hide future bands. show_all is trained-user override.
    return ("${show_all} = 'yes' or ${ca_band} >= %d" % k)

def ca_band_calc():
    parts, prev = [], None
    thresholds = [b[2] for b in BANDS[:-1]]
    expr = '18'
    for k in range(len(thresholds), 0, -1):
        expr = 'if(${ca_months_dec} < %s, %d, %s)' % (thresholds[k - 1], k, expr)
    return expr

row(type='calculate', name='ca_band', calculation=ca_band_calc())

for k, (key, label, end_m, weight) in enumerate(BANDS, start=1):
    items = ALL_BAND_ITEMS[k]
    first_no = items[0][0]
    row(type='begin group', name='g_' + key, label='%s (items %d-%d)' % (
        label, items[0][0], items[-1][0]), relevant=vb_expr(k))
    for no, text, speech in items:
        lbl = '%d. %s' % (no, text)
        if speech:
            lbl += ' [Speech/Language-related]'
        row(type='select_one dst_response', name='item_%d' % no, label=lbl,
            required='yes')
    row(type='end group')

    passes = ' + '.join("if(${item_%d} = 'pass', 1, 0)" % no for no, _, _ in items)
    nas = ' + '.join("if(${item_%d} = 'not_assessed', 1, 0)" % no for no, _, _ in items)
    fails = ' + '.join("if(${item_%d} = 'fail', 1, 0)" % no for no, _, _ in items)
    vis = '(%s)' % vb_expr(k)
    row(type='calculate', name='band_%d_passed_count' % k,
        label='NO_LABEL', calculation='if(%s, %s, 0)' % (vis, passes))
    row(type='calculate', name='band_%d_not_assessed' % k,
        label='NO_LABEL', calculation='if(%s, %s, 0)' % (vis, nas))
    row(type='calculate', name='band_%d_failed_count' % k,
        label='NO_LABEL', calculation='if(%s, %s, 0)' % (vis, fails))
    row(type='calculate', name='band_%d_earned_months' % k,
        label='NO_LABEL', calculation='${band_%d_passed_count} * (%s)' % (k, weight))

# ---------------------------------------------------------------- scoring
row(type='calculate', name='total_passed_items',
    calculation=' + '.join('${band_%d_passed_count}' % k for k in range(1, 19)))
row(type='calculate', name='total_failed_items',
    calculation=' + '.join('${band_%d_failed_count}' % k for k in range(1, 19)))
row(type='calculate', name='total_not_assessed_items',
    calculation=' + '.join('${band_%d_not_assessed}' % k for k in range(1, 19)))
row(type='calculate', name='developmental_age_months',
    calculation=' + '.join('${band_%d_earned_months}' % k for k in range(1, 19)))
row(type='calculate', name='developmental_age_years',
    calculation='floor(${developmental_age_months} div 12)')
row(type='calculate', name='developmental_age_remaining_months',
    calculation='round(${developmental_age_months} - (${developmental_age_years} * 12))')
row(type='calculate', name='score_status',
    calculation=("if(${age_valid} != 'yes', 'invalid_age', "
                 "if(${total_not_assessed_items} > 0, 'incomplete', 'complete'))"))
row(type='calculate', name='developmental_quotient_raw',
    calculation="if(${score_status} = 'complete' and ${ca_months_dec} > 0, "
                "${developmental_age_months} div ${ca_months_dec} * 100, '')")
row(type='calculate', name='developmental_quotient_rounded',
    calculation="if(${score_status} = 'complete', round(${developmental_quotient_raw}), '')")

row(type='text', name='na_reason',
    label='Reason item(s) were NOT ASSESSED (required when any applicable item is not assessed)',
    relevant='${total_not_assessed_items} > 0', required='yes')
row(type='text', name='clinical_notes', label='Clinical notes')

row(type='note', name='result_complete',
    label="**RESULTS**\nDevelopmental Age (DA): ${developmental_age_years}y "
          "${developmental_age_remaining_months}m (${developmental_age_months} months)\n"
          "Developmental Quotient (DQ): ${developmental_quotient_rounded} "
          "(raw: ${developmental_quotient_raw})\n"
          "DA = passed items weighted per age band. DQ = DA / chronological age x 100.\n"
          "This is a screening estimate and must be reviewed by a qualified clinician.",
    relevant="${score_status} = 'complete'")
row(type='note', name='result_incomplete',
    label='Score incomplete - clinician review required. Some applicable items were NOT ASSESSED; DA/DQ are withheld.',
    relevant="${score_status} = 'incomplete'")
row(type='note', name='result_invalid',
    label='Chronological age is outside the DST supported range (birth through 15 years). Result marked invalid_age.',
    relevant="${score_status} = 'invalid_age'")

# ---------------------------------------------------------------- Supporting Documents / Scale Upload
# Two fields: image for JPG/PNG, file for PDF — because single mediatype cannot reliably cover image/* and application/pdf in CHT 5.2.0 Enketo
row(type='begin group', name='g_supporting_docs', label='Supporting Documents / Scale Upload')
row(type='image', name='upload_scale_image', label='Upload completed scale / assessment sheet (image)',
    hint='Upload clear photos of the completed paper scale, consent form, or supporting assessment document. Accepts JPG, JPEG, PNG. For PDF use next field.',
    instance_type='binary')
row(type='file', name='upload_scale_pdf', label='Upload completed scale / assessment sheet (PDF)',
    hint='Upload a PDF of the completed paper scale if not already uploaded as image. Accepts PDF.',
    instance_type='binary')
row(type='text', name='upload_caption', label='Attachment caption / description (optional)')
row(type='end group')

ws_choices = wb.create_sheet('choices')
ws_choices.append(['list_name', 'name', 'label'])
for lst, rows in {
    'yesno': [('yes', 'Yes'), ('no', 'No')],
    'dst_response': [('pass', 'Pass'), ('fail', 'Fail'), ('not_assessed', 'Not assessed')],
}.items():
    for name, label in rows:
        ws_choices.append([lst, name, label])

ws_settings = wb.create_sheet('settings')
ws_settings.append(['form_title', 'form_id', 'version', 'default_language'])
ws_settings.append(['Developmental Screening Test (DST)', 'dst', '2026-08-22 01-00', 'en'])

wb.save(r'C:\Projects\EMRCHT\satoru-config\forms\app\dst.xlsx')
print('dst.xlsx created:', sum(len(v) for v in ALL_BAND_ITEMS.values()), 'items in',
      len(BANDS), 'bands')
