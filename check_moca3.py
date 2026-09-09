import openpyxl
wb = openpyxl.load_workbook(r'C:\Projects\EMRCHT\satoru-config\forms\app\moca_assessment.xlsx')
ws = wb['survey']

# Print all rows
print('=== All rows ===')
for row_idx in range(1, ws.max_row + 1):
    calc = ws.cell(row=row_idx, column=12).value
    name = ws.cell(row=row_idx, column=2).value
    type_ = ws.cell(row=row_idx, column=1).value
    print(f'Row {row_idx}: type={type_}, name={name}, calc={calc}')

print()
print('=== Max row:', ws.max_row)
print('=== Max col:', ws.max_column)