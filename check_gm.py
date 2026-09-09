import openpyxl
import re

wb = openpyxl.load_workbook(r'C:\Projects\EMRCHT\satoru-config\forms\app\ddst.xlsx')
ws = wb['survey']

# Print GM items relevant expressions
print('=== Gross Motor Items relevant expressions ===')
gm_items = []
for row_idx in range(129, 156):
    relevant = ws.cell(row=row_idx, column=7).value
    name = ws.cell(row=row_idx, column=3).value
    local_no = row_idx - 129 + 1  # GM local number 1 onwards
    if name or relevant:
        rel_str = str(relevant) if relevant else ''
        # Replace troublesome unicode chars
        for ch in ['\u25cb', '\u25a1']:
            rel_str = rel_str.replace(ch, '<square>')
        rel_str = rel_str[:40]
        # Extract age threshold from relevant expression
        age_match = re.search(r'>= ([\d.]+)', rel_str)
        age_threshold = float(age_match.group(1)) if age_match else None
        # Check if item is always visible (true()) or has age threshold
        always_visible = (rel_str == 'true()' or 'ca_months_dec' not in rel_str)
        gm_items.append((local_no, name, age_threshold, always_visible))
        print(f'GM{local_no}: relevant={rel_str}..., label={name}, age_thresh={age_threshold}, always_visible={always_visible}')

print()

# Check which GM items are visible at 9 months age (with show_all=no)
print('=== GM items visible at 9 months age (show_all=no) ===')
visible_at_9m = []
for local_no, name, age_threshold, always_visible in gm_items:
    if always_visible:
        visible_at_9m.append(local_no)
    elif age_threshold is not None and age_threshold <= 9.0:
        visible_at_9m.append(local_no)

visible_at_9m.sort()
print(f'GM items visible at 9 months: {visible_at_9m}')
print(f'Number of GM items visible at 9 months: {len(visible_at_9m)}')

# The expected items visible at 9 months should be GM1-GM10 (clean sequential list)
expected = list(range(1, 11))  # GM1 through GM10
actual = visible_at_9m
print(f'Expected visible: {expected}')
print(f'Actual visible: {actual}')

if actual == expected:
    print('\\nPASS: Clean sequential item list with no gaps at 9 months!')
else:
    print('\\nFAIL: Item list has gaps or missing items')
    missing = set(expected) - set(actual)
    extra = set(actual) - set(expected)
    if missing:
        print(f'Missing items: GM{list(missing)}')
    if extra:
        print(f'Extra items: GM{list(extra)}')
    
    # Show which specific items are missing/extra with their ages
    print('\\nDetails:')
    for local_no in range(1, 12):
        age = gm_items[local_no - 1][2] if local_no - 1 < len(gm_items) else None
        is_visible = local_no in visible_at_9m
        status = 'VISIBLE' if is_visible else 'HIDDEN'
        if age is not None:
            print(f'  GM{local_no}: age={age} months - {status}')
        else:
            print(f'  GM{local_no}: age=None (always visible) - {status}')